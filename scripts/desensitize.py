# -*- coding: utf-8 -*-
"""
面试人员名单脱敏：删除「姓名」「准考证号」两类个人身份列，保留岗位与分数信息。

本仓库不存储任何考生个人身份信息。原始名单由官方公开发布，但包含
姓名与准考证号，属于个人信息，因此入库前统一移除。

用法：
    python desensitize.py <输入目录> <输出目录>

输出文件名为原文件名去掉重复的 ".xlsx" 后缀。
"""
import os
import re
import sys

from openpyxl import load_workbook, Workbook

DROP_HEADERS = {'姓名', '准考证号'}


def norm(v):
    return re.sub(r'\s+', '', str(v)) if v is not None else ''


def desensitize(src_path, dst_path):
    wb = load_workbook(src_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    # 部分官方文件的 <dimension> 声明错误（如 A1:A1），只读模式会据此只读到 1 格，
    # 必须显式忽略声明的区域，改为真实扫描
    try:
        ws.reset_dimensions()
    except Exception:
        pass

    rows = ws.iter_rows(values_only=True)
    header = None
    for row in rows:
        if row and any(norm(c) for c in row):
            header = [norm(c) for c in row]
            break
    if header is None:
        raise RuntimeError('找不到表头: ' + src_path)

    keep = [i for i, h in enumerate(header) if h not in DROP_HEADERS]
    dropped = [header[i] for i in range(len(header)) if i not in keep]

    out = Workbook(write_only=True)
    sh = out.create_sheet('Sheet1')
    sh.append([header[i] for i in keep])
    n = 0
    for row in rows:
        if row is None or not any(c is not None and str(c).strip() for c in row):
            continue
        sh.append([row[i] if i < len(row) else None for i in keep])
        n += 1
    out.save(dst_path)
    wb.close()
    return len(header), dropped, n


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1
    src_dir, dst_dir = sys.argv[1], sys.argv[2]
    os.makedirs(dst_dir, exist_ok=True)
    files = sorted(f for f in os.listdir(src_dir)
                   if f.lower().endswith(('.xlsx', '.xls')) and not f.startswith('~$'))
    for f in files:
        src = os.path.join(src_dir, f)
        name = re.sub(r'(?i)(\.xlsx)+$', '.xlsx', f)
        dst = os.path.join(dst_dir, name)
        total, dropped, n = desensitize(src, dst)
        print('%-64s 列 %2d → 删除 %s / 保留 %d 行' % (
            f, total, '、'.join(dropped) if dropped else '（无）', n))
    print('\n完成，输出目录：' + dst_dir)
    return 0


if __name__ == '__main__':
    sys.exit(main())
